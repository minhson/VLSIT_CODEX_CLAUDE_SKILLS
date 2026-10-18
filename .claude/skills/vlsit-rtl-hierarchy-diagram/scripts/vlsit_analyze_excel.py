#!/usr/bin/env python3
"""Compact, read-only Excel/CSV analyzer for AI-assisted workbook updates.

Formats: .xlsx, .xlsm, .xltx, .xltm, .xls, .csv
Requires openpyxl. Legacy .xls conversion requires Microsoft Excel COM or LibreOffice.

Examples:
  python vlsit_analyze_excel.py book.xlsx
  python vlsit_analyze_excel.py book.xlsx --sheet MemoryMap --pretty
  python vlsit_analyze_excel.py book.xlsx --sheet MemoryMap --range A1:K64 --cells --pretty
  python vlsit_analyze_excel.py book.xlsx --sheet Booting --cells --output booting_cells.json
  python vlsit_analyze_excel.py book.xls --sheet Sheet1 --cells
  python vlsit_analyze_excel.py data.csv --cells --range A1:F200
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import os
import platform
import shutil
import subprocess
import tempfile
import zipfile
from collections import Counter
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, range_boundaries
except ImportError as exc:
    raise SystemExit("Missing dependency: openpyxl. Install it with: python -m pip install openpyxl") from exc

VERSION = "2.0"
EXCEL_FORMATS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
SUPPORTED = EXCEL_FORMATS | {".xls", ".xlsb", ".csv"}


class AnalysisError(Exception):
    pass


def json_value(value: Any) -> Any:
    """Convert a cell value to JSON without changing its text or whitespace."""
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    if hasattr(value, "text") and isinstance(value.text, str):
        return value.text
    return str(value)


def display_value(value: Any, limit: int) -> Any:
    """Shorten only preview text; cell dumps always use json_value unchanged."""
    result = json_value(value)
    if isinstance(result, str) and len(result) > limit:
        return result[:limit] if limit < 4 else result[: limit - 3] + "..."
    return result


def color_text(color: Any) -> str | None:
    if color is None:
        return None
    if color.type == "rgb":
        return str(color.rgb)
    if color.type == "theme":
        return f"theme:{color.theme}"
    if color.type == "indexed":
        return f"indexed:{color.indexed}"
    return str(color.type) if color.type else None


def style_summary(cells: list[Any], cap: int = 8) -> dict[str, Any]:
    fonts: Counter[str] = Counter()
    fills: Counter[str] = Counter()
    number_formats: Counter[str] = Counter()
    style_ids: Counter[int] = Counter()
    for cell in cells:
        f = cell.font
        font_key = " ".join(part for part in (
            str(f.name or "default"), str(f.sz or ""),
            "bold" if f.bold else "", "italic" if f.italic else "",
            f"color={color_text(f.color)}" if f.color else "",
        ) if part)
        fonts[font_key] += 1
        fill = cell.fill
        fills[f"{fill.fill_type or 'none'}:{color_text(fill.fgColor) or 'default'}"] += 1
        number_formats[cell.number_format or "General"] += 1
        style_ids[cell.style_id] += 1
    return {
        "fonts": [{"style": key, "cells": count} for key, count in fonts.most_common(cap)],
        "fills": [{"style": key, "cells": count} for key, count in fills.most_common(cap)],
        "number_formats": [{"format": key, "cells": count} for key, count in number_formats.most_common(cap)],
        "style_ids": [{"id": key, "cells": count} for key, count in style_ids.most_common(cap)],
    }


def make_preview(ws: Any, rows: list[int], min_col: int, max_col: int,
                 max_cols: int, max_text: int) -> list[dict[str, Any]]:
    end_col = min(max_col, min_col + max_cols - 1)
    result = []
    for row_num in rows:
        values = [display_value(ws.cell(row_num, col).value, max_text)
                  for col in range(min_col, end_col + 1)]
        result.append({
            "row": row_num,
            "from": get_column_letter(min_col),
            "values": values,
            **({"more_columns": max_col - end_col} if end_col < max_col else {}),
        })
    return result


def cell_dump(ws: Any, cached_ws: Any, args: argparse.Namespace,
              bounds: tuple[int, int, int, int] | None) -> list[dict[str, Any]]:
    min_col, min_row, max_col, max_row = bounds or (1, 1, ws.max_column, ws.max_row)
    cells = []
    for cell in sorted(ws._cells.values(), key=lambda item: (item.row, item.column)):
        inside = min_row <= cell.row <= max_row and min_col <= cell.column <= max_col
        if not inside:
            continue
        if cell.value is None and not (args.include_styled_blanks and cell.has_style):
            continue
        record = {
            "a": cell.coordinate,
            "kind": "formula" if cell.data_type == "f" else ("error" if cell.data_type == "e" else "value"),
            "excel_type": cell.data_type,
            "value": json_value(cell.value),
            "style_id": cell.style_id,
        }
        if cell.data_type == "f":
            cached = cached_ws.cell(cell.row, cell.column).value if cached_ws else None
            record["cached_value"] = json_value(cached)
        if cell.number_format and cell.number_format != "General":
            record["number_format"] = cell.number_format
        if cell.hyperlink:
            record["hyperlink"] = {
                "target": cell.hyperlink.target,
                "location": cell.hyperlink.location,
            }
        if cell.comment:
            record["comment"] = json_value(cell.comment.text)
        cells.append(record)
        if len(cells) > args.max_dump_cells:
            raise AnalysisError(
                f"Cell dump exceeds --max-dump-cells={args.max_dump_cells}; narrow it with --range or raise the limit."
            )
    return cells


def analyze_sheet(ws: Any, cached_ws: Any, args: argparse.Namespace,
                  detail: bool, include_cells: bool) -> dict[str, Any]:
    populated = [cell for cell in ws._cells.values() if cell.value is not None]
    if populated:
        min_row = min(c.row for c in populated)
        max_row = max(c.row for c in populated)
        min_col = min(c.column for c in populated)
        max_col = max(c.column for c in populated)
        content_ref = f"{get_column_letter(min_col)}{min_row}:{get_column_letter(max_col)}{max_row}"
        nonempty_rows = sorted({c.row for c in populated})
    else:
        min_row = max_row = min_col = max_col = 1
        content_ref = None
        nonempty_rows = []

    bounds = None
    if args.range:
        bounds = range_boundaries(args.range)
        if None in bounds:
            raise AnalysisError(f"Range must be bounded, for example A1:K64: {args.range}")
        min_col, min_row, max_col, max_row = bounds
        nonempty_rows = sorted({c.row for c in populated if min_row <= c.row <= max_row})

    sample_rows = nonempty_rows[: args.preview_rows]
    if nonempty_rows and nonempty_rows[-1] not in sample_rows:
        sample_rows.append(nonempty_rows[-1])

    scoped = [c for c in populated if bounds is None or
              (min_row <= c.row <= max_row and min_col <= c.column <= max_col)]
    formulas = [c for c in populated if c.data_type == "f"]
    errors = [c for c in populated if c.data_type == "e"]
    links = [c for c in populated if c.hyperlink is not None]
    comments = [c for c in populated if c.comment is not None]

    formula_samples = []
    if detail:
        for cell in formulas[: args.sample_limit]:
            cached = cached_ws.cell(cell.row, cell.column).value if cached_ws else None
            formula_samples.append({
                "cell": cell.coordinate,
                "formula": display_value(cell.value, args.max_text),
                "cached_value": json_value(cached),
            })

    tables = []
    for table in ws.tables.values():
        tables.append({
            "name": table.name,
            "range": table.ref,
            "filter": bool(table.autoFilter is not None),
            "columns": [col.name for col in (table.tableColumns or [])[: max(args.sample_limit, 20)]],
        })

    widths = []
    for col_letter, dimension in ws.column_dimensions.items():
        if dimension.width is not None or dimension.hidden:
            first_col = dimension.min or ws[f"{col_letter}1"].column
            last_col = dimension.max or first_col
            col_range = (get_column_letter(first_col) if first_col == last_col else
                         f"{get_column_letter(first_col)}:{get_column_letter(last_col)}")
            widths.append({"columns": col_range,
                           "width": round(dimension.width, 2) if dimension.width is not None else None,
                           **({"hidden": True} if dimension.hidden else {})})
            if detail and len(widths) >= args.sample_limit:
                break

    formula_cache_cells = 0
    cached_formula_errors = 0
    if cached_ws:
        for cell in formulas:
            cached_cell = cached_ws.cell(cell.row, cell.column)
            if cached_cell.value is not None:
                formula_cache_cells += 1
                if cached_cell.data_type == "e":
                    cached_formula_errors += 1

    validation = getattr(ws, "data_validations", None)
    report = {
        "name": ws.title,
        "state": ws.sheet_state,
        "declared_range": ws.calculate_dimension(),
        "content_range": content_ref,
        "nonempty_cells": len(populated),
        "formulas": len(formulas),
        "formula_cache_cells": formula_cache_cells,
        "cached_formula_errors": cached_formula_errors,
        "formula_samples": formula_samples,
        "error_cells": len(errors),
        "error_samples": [{"cell": c.coordinate, "value": json_value(c.value)}
                          for c in errors[: args.sample_limit]] if detail else [],
        "hyperlinks": {"count": len(links), "samples": [
            {"cell": c.coordinate, "text": display_value(c.value, args.max_text),
             "target": c.hyperlink.target, "location": c.hyperlink.location}
            for c in links[: args.sample_limit]
        ] if detail else []},
        "comments": {"count": len(comments), "samples": [
            {"cell": c.coordinate, "text": display_value(c.comment.text, args.max_text)}
            for c in comments[: args.sample_limit]
        ] if detail else []},
        "tables": tables,
        "autofilter": ws.auto_filter.ref,
        "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
        "merged_ranges": {"count": len(ws.merged_cells.ranges),
                           "samples": [str(r) for r in list(ws.merged_cells.ranges)[: args.sample_limit] if detail]},
        "validations": len(validation.dataValidation) if validation else 0,
        "images": len(getattr(ws, "_images", [])),
        "charts": len(getattr(ws, "_charts", [])),
        "hidden_rows": [r for r, dim in ws.row_dimensions.items() if dim.hidden][: args.sample_limit] if detail else [],
        "hidden_columns": [w["columns"] for w in widths if w.get("hidden")],
        "column_widths": widths if detail else [],
        "print_area": str(ws.print_area) if ws.print_area else None,
        "style_summary": style_summary(populated) if detail else {
            "top_fonts": style_summary(populated, cap=3)["fonts"]
        },
        "preview": make_preview(ws, sample_rows, min_col, max_col,
                                 args.preview_cols, args.max_text) if detail and content_ref else [],
    }
    if include_cells:
        report["cell_dump"] = cell_dump(ws, cached_ws, args, bounds)
        report["cell_dump_scope"] = args.range or content_ref or "entire_sheet"
    return report


def find_soffice(requested: str | None) -> str | None:
    candidates = [requested, os.environ.get("SOFFICE_PATH"), shutil.which("soffice"),
                  shutil.which("libreoffice")]
    if platform.system() == "Windows":
        candidates.extend([
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        ])
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(Path(candidate))
        if candidate and shutil.which(candidate):
            return shutil.which(candidate)
    return None


def convert_xls_with_excel(source: Path, destination: Path, temp_dir: Path) -> tuple[bool, str]:
    if platform.system() != "Windows":
        return False, "Excel COM is available only on Windows."
    powershell = shutil.which("powershell.exe") or shutil.which("powershell")
    if not powershell:
        return False, "PowerShell was not found."
    ps_file = temp_dir / "convert_xls.ps1"
    ps_file.write_text(r'''
$ErrorActionPreference = "Stop"
$src = $args[0]
$dst = $args[1]
$excel = $null
$book = $null
try {
    $excel = New-Object -ComObject Excel.Application
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $excel.AutomationSecurity = 3
    $excel.Calculation = -4135
    $book = $excel.Workbooks.Open($src, 0, $true)
    $book.CheckCompatibility = $false
    $book.SaveAs($dst, 51)
    $book.Close($false)
    $excel.Quit()
    Write-Output "CONVERTED"
} finally {
    if ($book) { try { $book.Close($false) } catch {} }
    if ($excel) { try { $excel.Quit() } catch {}; try { [void][System.Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel) } catch {} }
}
''', encoding="utf-8")
    try:
        result = subprocess.run([powershell, "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                                 "-File", str(ps_file), str(source), str(destination)],
                                capture_output=True, text=True, timeout=120)
    except Exception as exc:
        return False, f"Excel COM launch failed: {exc}"
    if result.returncode == 0 and destination.is_file():
        return True, "Excel COM converted the source to a temporary XLSX; source was opened read-only."
    detail = (result.stderr or result.stdout or "Excel COM conversion did not create output.").strip()
    return False, detail[-1000:]


def convert_legacy_excel(source: Path, args: argparse.Namespace, temp_dir: Path) -> tuple[Path, str, str]:
    destination = temp_dir / "converted" / (source.stem + ".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)
    ok, message = convert_xls_with_excel(source, destination, temp_dir)
    if ok:
        return destination, "excel_com", message

    soffice = find_soffice(args.soffice)
    if source.suffix.lower() == ".xls" and soffice:
        profile = temp_dir / "lo_profile"
        command = [soffice, "--headless", "--nologo", "--nodefault", "--nolockcheck",
                   f"-env:UserInstallation={profile.as_uri()}", "--convert-to", "xlsx",
                   "--outdir", str(destination.parent), str(source)]
        try:
            result = subprocess.run(command, capture_output=True, text=True, timeout=180)
        except Exception as exc:
            raise AnalysisError(f"LibreOffice conversion failed: {exc}") from exc
        if result.returncode == 0 and destination.is_file():
            return destination, "libreoffice", (
                "LibreOffice converted the source to a temporary XLSX; source was not changed. "
                "Cached formula results may be recalculated during conversion."
            )
        detail = (result.stderr or result.stdout or "Conversion did not create output.").strip()
        raise AnalysisError(f"Could not convert .xls with LibreOffice: {detail[-1000:]}")

    if source.suffix.lower() == ".xlsb":
        raise AnalysisError(
            ".xlsb requires Microsoft Excel automation to preserve formulas and cached values. "
            "Excel COM was unavailable; the source workbook was not changed. "
            f"Excel COM detail: {message}"
        )
    raise AnalysisError(
        ".xls requires Microsoft Excel (Windows COM) or LibreOffice to preserve formulas and cached values. "
        "Neither converter was available. Install one or pass --soffice <path>; the source workbook was not changed. "
        f"Excel COM detail: {message}"
    )


def make_temp_context(source_dir: Path) -> tempfile.TemporaryDirectory:
    """Use a temporary folder beside the workbook; fall back to the system temp folder."""
    try:
        return tempfile.TemporaryDirectory(prefix=".excel_analyzer_", dir=str(source_dir))
    except OSError:
        try:
            return tempfile.TemporaryDirectory(prefix="excel_analyzer_")
        except OSError as exc:
            raise AnalysisError(f"Could not create a temporary folder for read-only .xls conversion: {exc}") from exc


def analyze_csv(path: Path, args: argparse.Namespace) -> dict[str, Any]:
    raw = path.read_bytes()
    encoding = "utf-8-sig"
    text = None
    for candidate in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise AnalysisError("Could not decode CSV text.")

    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        dialect = csv.excel
        delimiter = ","

    rows = list(csv.reader(io.StringIO(text, newline=""), dialect))
    if args.sheets and args.sheets != ["CSV"]:
        raise AnalysisError("CSV has one virtual sheet named 'CSV'; use --sheet CSV or omit --sheet.")
    if args.range:
        bounds = range_boundaries(args.range)
        if None in bounds:
            raise AnalysisError(f"Range must be bounded, for example A1:F20: {args.range}")
    else:
        bounds = None
    min_col, min_row, max_col, max_row = bounds or (1, 1, max((len(r) for r in rows), default=1), max(1, len(rows)))

    widths = [len(row) for row in rows]
    max_cols = max(1, max(widths, default=0))
    nonempty = [(rno, cno, value) for rno, row in enumerate(rows, 1)
                for cno, value in enumerate(row, 1) if value != ""]
    if nonempty:
        content_range = (f"{get_column_letter(min(c for _, c, _ in nonempty))}"
                         f"{min(r for r, _, _ in nonempty)}:"
                         f"{get_column_letter(max(c for _, c, _ in nonempty))}"
                         f"{max(r for r, _, _ in nonempty)}")
    else:
        content_range = None
    formula_like = sum(1 for _, _, value in nonempty if value.startswith("="))
    search_hits = []
    if args.search:
        needle = args.search.casefold()
        for rno, cno, value in nonempty:
            if needle in value.casefold():
                search_hits.append({"cell": f"{get_column_letter(cno)}{rno}", "text": display_value(value, args.max_text)})
                if len(search_hits) >= args.sample_limit:
                    break

    detail = bool(args.sheets or args.range or args.detail)
    preview = []
    if detail:
        row_ids = list(range(max(1, min_row), min(len(rows), max_row) + 1))
        if not args.range:
            row_ids = row_ids[: args.preview_rows]
            if rows and len(rows) not in row_ids:
                row_ids.append(len(rows))
        elif len(row_ids) > args.preview_rows + 1:
            row_ids = row_ids[: args.preview_rows] + [row_ids[-1]]
        for row_num in row_ids:
            row = rows[row_num - 1]
            values = [display_value(row[c - 1], args.max_text) if c <= len(row) else None
                      for c in range(min_col, min(max_col, min_col + args.preview_cols - 1) + 1)]
            preview.append({"row": row_num, "from": get_column_letter(min_col), "values": values})

    cell_records = []
    if args.cells:
        selected = [(rno, cno, value) for rno, row in enumerate(rows, 1)
                    for cno, value in enumerate(row, 1)
                    if min_row <= rno <= max_row and min_col <= cno <= max_col]
        if len(selected) > args.max_dump_cells:
            raise AnalysisError(
                f"CSV cell dump exceeds --max-dump-cells={args.max_dump_cells}; narrow it with --range or raise the limit."
            )
        cell_records = [{"a": f"{get_column_letter(cno)}{rno}", "kind": "text",
                         "value": value, "starts_with_equal": value.startswith("="),
                         "formula_semantics": "not_stored_in_csv"}
                        for rno, cno, value in selected]

    report = {
        "analyzer": VERSION,
        "file": str(path),
        "format": ".csv",
        "size_bytes": len(raw),
        "report_mode": "detailed" if detail else "workbook_summary",
        "sheet_count": 1,
        "sheet_names": ["CSV"],
        "csv": {"encoding": encoding, "delimiter": "TAB" if delimiter == "\t" else delimiter,
                "rows": len(rows), "max_columns": max_cols,
                "ragged_rows": sum(1 for width in widths if width != max_cols),
                "nonempty_cells": len(nonempty), "formula_like_text_cells": formula_like,
                "formula_warning": "CSV stores text only; cells beginning with '=' are not reliably distinguishable as formulas."},
        "sheets": [{
            "name": "CSV", "state": "visible", "declared_range": f"A1:{get_column_letter(max_cols or 1)}{len(rows) or 1}",
            "content_range": content_range, "nonempty_cells": len(nonempty),
            "formulas": None, "formula_cache_cells": None, "formula_note": "CSV has no formula metadata or cached formula values.",
            "error_cells": None, "hyperlinks": {"count": 0, "samples": []}, "tables": [],
            "autofilter": None, "freeze_panes": None, "merged_ranges": {"count": 0, "samples": []},
            "style_summary": None, "preview": preview,
            **({"cell_dump": cell_records, "cell_dump_scope": args.range or content_range or "entire_csv"} if args.cells else {}),
        }],
        **({"search": {"term": args.search, "sampled_hits": search_hits,
                        "limit_reached": len(search_hits) >= args.sample_limit}} if args.search else {}),
    }
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="Read-only analyzer for Excel workbooks and CSV files.")
    parser.add_argument("file", type=Path, help="Workbook or CSV path")
    parser.add_argument("--sheet", action="append", dest="sheets", help="Limit to a sheet; repeat to select multiple")
    parser.add_argument("--range", help="Bounded cell range, for example A1:K64")
    parser.add_argument("--search", help="Find text in cell values or formulas")
    parser.add_argument("--cells", action="store_true", help="Dump every in-scope cell with address, raw value/formula, style ID, and formula cache")
    parser.add_argument("--max-dump-cells", type=int, default=10000, help="Safety limit for --cells (default: 10000)")
    parser.add_argument("--include-styled-blanks", action="store_true", help="Include formatted empty cells in a bounded --cells range")
    parser.add_argument("--preview-rows", type=int, default=4, help="Preview populated rows (default: 4)")
    parser.add_argument("--preview-cols", type=int, default=12, help="Maximum preview columns (default: 12)")
    parser.add_argument("--sample-limit", type=int, default=8, help="Maximum examples per feature (default: 8)")
    parser.add_argument("--max-text", type=int, default=96, help="Maximum preview text length; cell dumps stay exact (default: 96)")
    parser.add_argument("--soffice", help="Path to LibreOffice soffice executable for legacy .xls conversion")
    parser.add_argument("--detail", action="store_true", help="Include previews and style examples for all sheets")
    parser.add_argument("--pretty", action="store_true", help="Indent JSON; compact JSON is the default")
    parser.add_argument("--output", type=Path, help="Write report to a JSON file instead of stdout")
    args = parser.parse_args()

    path = args.file.expanduser().resolve()
    if not path.exists():
        parser.error(f"File not found: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        parser.error(f"Unsupported format {suffix!r}; supported: {', '.join(sorted(SUPPORTED))}")
    if args.output and args.output.expanduser().resolve() == path:
        parser.error("Report output cannot overwrite the source file.")
    if args.range and suffix != ".csv" and (not args.sheets or len(args.sheets) != 1):
        parser.error("--range requires exactly one --sheet for Excel workbooks.")
    if args.cells and suffix != ".csv" and len(args.sheets or []) != 1:
        parser.error("For safety, --cells on an Excel workbook requires exactly one --sheet.")
    if args.include_styled_blanks and (not args.cells or not args.range or suffix == ".csv"):
        parser.error("--include-styled-blanks requires --cells and a bounded Excel --range.")
    if args.max_dump_cells < 1 or min(args.preview_rows, args.preview_cols, args.sample_limit, args.max_text) < 1:
        parser.error("limits must be positive")

    try:
        if suffix == ".csv":
            if args.cells and args.sheets and args.sheets != ["CSV"]:
                raise AnalysisError("CSV has one virtual sheet named 'CSV'; use --sheet CSV or omit --sheet.")
            report = analyze_csv(path, args)
        else:
            temp_context = make_temp_context(path.parent) if suffix in {".xls", ".xlsb"} else None
            try:
                if suffix in {".xls", ".xlsb"}:
                    effective_path, backend, conversion_note = convert_legacy_excel(
                        path, args, Path(temp_context.name)
                    )
                else:
                    effective_path, backend, conversion_note = path, "openpyxl", None

                with zipfile.ZipFile(effective_path) as zf:
                    corrupt_part = zf.testzip()
                    zip_entries = len(zf.infolist())
                if corrupt_part:
                    raise AnalysisError(f"Corrupt ZIP member in analysis copy: {corrupt_part}")

                keep_vba = suffix in {".xlsm", ".xltm"}
                try:
                    wb = load_workbook(effective_path, data_only=False, read_only=False,
                                       keep_vba=keep_vba, keep_links=True)
                    cached_wb = load_workbook(effective_path, data_only=True, read_only=False,
                                              keep_vba=keep_vba, keep_links=True)
                except Exception as exc:
                    raise AnalysisError(f"Could not read workbook: {type(exc).__name__}: {exc}") from exc

                all_names = wb.sheetnames
                chosen = args.sheets or all_names
                missing = [name for name in chosen if name not in all_names]
                if missing:
                    raise AnalysisError(f"Unknown sheet(s): {', '.join(missing)}. Available: {', '.join(all_names)}")
                detail = bool(args.sheets or args.range or args.detail or args.cells)
                report = {
                    "analyzer": VERSION,
                    "file": str(path),
                    "format": suffix,
                    "analysis_backend": backend,
                    "conversion_note": conversion_note,
                    "size_bytes": path.stat().st_size,
                    "zip_integrity": "ok",
                    "zip_entries": zip_entries,
                    "macro_container": suffix in {".xlsm", ".xltm"},
                    "legacy_format_may_contain_macros": suffix in {".xls", ".xlsb"},
                    "macro_code_analyzed": False,
                    "sheet_count": len(all_names),
                    "sheet_names": all_names,
                    "defined_name_count": len(wb.defined_names),
                    "defined_names": list(wb.defined_names.keys())[: args.sample_limit],
                    "external_link_count": len(getattr(wb, "_external_links", [])),
                    "formula_values_note": "Formula text is preserved; cached_value is reported separately. This script does not recalculate formulas.",
                    "report_mode": "detailed" if detail else "workbook_summary",
                    "sheets": [analyze_sheet(wb[name], cached_wb[name], args, detail, args.cells)
                               for name in chosen],
                }
                if args.search:
                    needle = args.search.casefold()
                    hits = []
                    for name in chosen:
                        ws = wb[name]
                        for cell in ws._cells.values():
                            if cell.value is not None and needle in str(cell.value).casefold():
                                hits.append({"sheet": name, "cell": cell.coordinate,
                                             "value": display_value(cell.value, args.max_text)})
                                if len(hits) >= args.sample_limit:
                                    break
                        if len(hits) >= args.sample_limit:
                            break
                    report["search"] = {"term": args.search, "sampled_hits": hits,
                                         "limit_reached": len(hits) >= args.sample_limit}
                wb.close()
                cached_wb.close()
            finally:
                if temp_context:
                    temp_context.cleanup()

        output_text = json.dumps(report, ensure_ascii=not bool(args.output), indent=2 if args.pretty else None,
                                 separators=None if args.pretty else (",", ":"))
        if args.output:
            args.output.expanduser().resolve().write_text(output_text + "\n", encoding="utf-8")
        else:
            print(output_text)
        return 0
    except (AnalysisError, OSError, zipfile.BadZipFile, csv.Error) as exc:
        parser.error(str(exc))
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
